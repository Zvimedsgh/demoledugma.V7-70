import openpyxl

filepath = r"C:\ledugma\DEMO\Demo_Reports_Syatem_V7.70.xlsm"
wb = openpyxl.load_workbook(filepath, data_only=True)
ws = wb.worksheets[0]

print(f"Number of tables in Management sheet: {len(ws.tables)}")
for table in ws.tables.values():
    print(f"Table name: {table.name}, Ref: {table.ref}")

