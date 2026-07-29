import openpyxl

filepath = r"C:\ledugma\DEMO\Demo_Reports_Syatem_V7.70.xlsm"
wb = openpyxl.load_workbook(filepath, data_only=True)

sheet_name = "רשימות"
if sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for col in range(1, 6):
        last_row = 1
        for r in range(1, ws.max_row + 1):
            if ws.cell(row=r, column=col).value is not None:
                last_row = r
        print(f"Col {col} max row: {last_row}")
else:
    print("Sheet not found")
