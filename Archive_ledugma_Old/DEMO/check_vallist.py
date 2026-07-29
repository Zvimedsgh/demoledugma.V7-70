import openpyxl

filepath = r"C:\ledugma\DEMO\Demo_Reports_Syatem_V7.70.xlsm"
wb = openpyxl.load_workbook(filepath, data_only=True)

sheet_name = "ניהול"
if sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
else:
    ws = wb.worksheets[0]

# Collect main branches
main_branches = set()
for row in range(40, 130):
    valB = ws.cell(row=row, column=2).value
    if valB:
        main_branches.add(str(valB).strip())

valList = ",".join(main_branches)
print(f"Number of unique main branches: {len(main_branches)}")
print(f"Length of comma-separated valList for main branches: {len(valList)}")

# Collect branches
branches = set()
for row in range(40, 130):
    valA = ws.cell(row=row, column=1).value
    if valA:
        branches.add(str(valA).strip())

valList2 = ",".join(branches)
print(f"Number of unique branches: {len(branches)}")
print(f"Length of comma-separated valList for branches: {len(valList2)}")
