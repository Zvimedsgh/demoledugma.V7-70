import openpyxl
import io

filepath = r"C:\ledugma\DEMO\Demo_Reports_Syatem_V7.70.xlsm"
try:
    wb = openpyxl.load_workbook(filepath, data_only=True)
    
    # Try to find the management sheet
    sheet_name = "ניהול"
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.worksheets[0]
        
    with io.open("mgmt_output2.txt", "w", encoding="utf-8") as f:
        f.write(f"Sheet name: {ws.title}\n")
        
        for row in range(40, 130):
            valA = ws.cell(row=row, column=1).value
            valB = ws.cell(row=row, column=2).value
            if valA is not None or valB is not None:
                if row >= 98 and row <= 105:
                    f.write(f"Row {row}: A={repr(valA)}, B={repr(valB)}\n")
                if row == 128:
                    f.write(f"Row {row}: A={repr(valA)}, B={repr(valB)}\n")
        
        last_row = 1
        for r in range(1, ws.max_row + 1):
            if ws.cell(row=r, column=1).value is not None:
                last_row = r
        f.write(f"Max row in Column A: {last_row}\n")
    
except Exception as e:
    with io.open("mgmt_output2.txt", "w", encoding="utf-8") as f:
        f.write(f"Error: {e}\n")
