import openpyxl

wb = openpyxl.load_workbook(r'C:\ledugma\DEMO\Demo_Reports_Syatem_V7.70.xlsm', data_only=True)
ws = wb['ראשי'] if 'ראשי' in wb.sheetnames else wb.worksheets[0]

with open('cells.txt', 'w', encoding='utf-8') as f:
    f.write(f"F2={ws['F2'].value}\n")
    f.write(f"F3={ws['F3'].value}\n")
    f.write(f"F4={ws['F4'].value}\n")
    f.write(f"A1={ws['A1'].value}\n")
    f.write(f"A2={ws['A2'].value}\n")
