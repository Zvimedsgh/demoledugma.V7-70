import openpyxl

filepath = r"C:\ledugma\DEMO\Demo_Reports_Syatem_V7.70.xlsm"
try:
    wb = openpyxl.load_workbook(filepath, data_only=True)
    
    # Try to find the management sheet
    sheet_name = "ניהול"
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        # Just grab the first sheet if we can't find it
        ws = wb.worksheets[0]
        
    print(f"Sheet name: {ws.title}")
    
    for row in range(40, 130):
        valA = ws.cell(row=row, column=1).value
        valB = ws.cell(row=row, column=2).value
        if valA is not None or valB is not None:
            if row >= 98 and row <= 105:
                print(f"Row {row}: A={repr(valA)}, B={repr(valB)}")
            if row == 128:
                print(f"Row {row}: A={repr(valA)}, B={repr(valB)}")
    
    # also print the last row with data in column A
    last_row = 1
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value is not None:
            last_row = r
    print(f"Max row in Column A: {last_row}")
    
except Exception as e:
    print(f"Error: {e}")
