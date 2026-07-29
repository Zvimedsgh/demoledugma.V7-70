import openpyxl

wb = openpyxl.load_workbook(r'C:\ledugma\DEMO\Demo_Reports_Syatem_V7.70.xlsm', data_only=True)
ws = wb['הגדרות'] if 'הגדרות' in wb.sheetnames else wb['NIHUL']
val = ws['K46'].value

with open('k46.txt', 'w', encoding='utf-8') as f:
    f.write(f"K46={val}")
