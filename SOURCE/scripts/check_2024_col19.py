import sys
import openpyxl

file_path = r'c:\LEVAV PROJECT\SOURCE\2024.xlsx'
try:
    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    sheet = wb.active
    
    print("Sampling column 19 (S) from 2024.xlsx:")
    for i, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, min_col=19, max_col=19, values_only=True)):
        print(f"Row {i+1}: {row[0]} (Type: {type(row[0]).__name__})")
        
    print("\nChecking for any values that are NOT dates or look like weird numbers...")
    invalid_count = 0
    total_count = 0
    for row in sheet.iter_rows(min_row=2, min_col=19, max_col=19, values_only=True):
        val = row[0]
        if val is not None:
            total_count += 1
            # In Excel, dates might come out as datetime objects, strings, or numbers
            if type(val).__name__ not in ('datetime', 'time', 'date'):
                # If it's a string, does it look like a date?
                sval = str(val).strip()
                if sval == "" or sval == "0":
                    pass
                elif sval.replace('.', '').replace('-', '').replace('/', '').isdigit():
                    # It's a number. Is it a reasonable Excel date serial? (e.g. 40000 - 50000)
                    try:
                        num = float(sval)
                        if num < 30000 or num > 60000:
                            if invalid_count < 10:
                                print(f"Weird numeric value found: {val}")
                            invalid_count += 1
                    except:
                        pass
                else:
                    if invalid_count < 10:
                        print(f"Non-date string found: {val}")
                    invalid_count += 1
                    
    print(f"\nTotal rows checked: {total_count}")
    print(f"Total invalid/weird values found: {invalid_count}")
except Exception as e:
    print(f"Error reading Excel file: {e}")
